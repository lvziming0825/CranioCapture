"""
CranioCapture

Clinical measurement and metric calibration GUI.

This tool provides:

1. Interactive scale calibration
2. Real-world distance measurement
3. Excel export of measurement results
4. Reproducible metric calibration workflow

Outputs:
- Calibration coefficient
- Distance measurements (mm)
- Excel report

Author:
Richard LYU
"""


import cv2
import numpy as np
import math
import os
import tkinter as tk
from tkinter import filedialog, simpledialog, messagebox
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side


root = tk.Tk()
root.withdraw()

IMAGE_PATH = filedialog.askopenfilename(
    title="Select image",
    filetypes=[("Image Files", "*.png *.jpg *.jpeg *.bmp *.tif *.tiff")]
)

if not IMAGE_PATH:
    raise Exception("No image selected.")

OUTPUT_DIR = filedialog.askdirectory(title="Select output folder")

if not OUTPUT_DIR:
    raise Exception("No output folder selected.")


img = cv2.imread(IMAGE_PATH)

if img is None:
    raise FileNotFoundError(f"Cannot read image: {IMAGE_PATH}")

img_h, img_w = img.shape[:2]
image_name = os.path.splitext(os.path.basename(IMAGE_PATH))[0]

excel_path = os.path.join(
    OUTPUT_DIR,
    f"{image_name}_measurement_result.xlsx"
)


WINDOW_NAME = "CranioCapture Clinical Measurement"

canvas_w = 1200
canvas_h = 900

scale = 1.0
offset_x = 0
offset_y = 0

dragging = False
last_mouse = None

mode = "calibration"

current_points = []
calibration_records = []
measurement_records = []

final_mm_per_px = None

pending_calibration = False
pending_measurement = False


def distance(p1, p2):
    return math.sqrt(
        (p1[0] - p2[0]) ** 2 +
        (p1[1] - p2[1]) ** 2
    )


def fit_image_to_window():
    global scale, offset_x, offset_y

    margin = 40
    scale_x = (canvas_w - margin) / img_w
    scale_y = (canvas_h - margin) / img_h
    scale = min(scale_x, scale_y)

    display_w = int(img_w * scale)
    display_h = int(img_h * scale)

    offset_x = int((canvas_w - display_w) / 2)
    offset_y = int((canvas_h - display_h) / 2)


def image_to_screen(p):
    return int(p[0] * scale + offset_x), int(p[1] * scale + offset_y)


def screen_to_image(x, y):
    return int((x - offset_x) / scale), int((y - offset_y) / scale)


def midpoint(p1, p2):
    return int((p1[0] + p2[0]) / 2), int((p1[1] + p2[1]) / 2)


def save_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "Measurement Results"

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    sub_fill = PatternFill("solid", fgColor="E2F0D9")
    thin = Side(style="thin", color="AAAAAA")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws["A1"] = "Interactive Clinical Measurement Result"
    ws["A1"].font = Font(bold=True, size=15)

    info = [
        ["Image Name", image_name],
        ["Image Path", IMAGE_PATH],
        ["Output Excel", excel_path],
        ["Date Time", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
        ["Final mean mm/px", final_mm_per_px if final_mm_per_px else ""],
    ]

    row = 3
    for k, v in info:
        ws.cell(row=row, column=1, value=k)
        ws.cell(row=row, column=2, value=v)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Calibration Records")
    ws.cell(row=row, column=1).font = Font(bold=True)
    row += 1

    headers = [
        "Index", "P1_x", "P1_y", "P2_x", "P2_y",
        "Pixel Distance", "Real Distance mm", "mm_per_px"
    ]

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center")

    row += 1

    for r in calibration_records:
        values = [
            r["index"], r["p1"][0], r["p1"][1], r["p2"][0], r["p2"][1],
            r["pixel_distance"], r["real_mm"], r["mm_per_px"]
        ]

        for col, v in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.border = border
            cell.alignment = Alignment(horizontal="center")

        row += 1

    row += 2
    ws.cell(row=row, column=1, value="Measurement Records")
    ws.cell(row=row, column=1).font = Font(bold=True)
    row += 1

    headers = [
        "Index", "P1_x", "P1_y", "P2_x", "P2_y",
        "Pixel Distance", "Distance mm"
    ]

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = Font(bold=True)
        cell.fill = sub_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center")

    row += 1

    for r in measurement_records:
        values = [
            r["index"], r["p1"][0], r["p1"][1], r["p2"][0], r["p2"][1],
            r["pixel_distance"], r["distance_mm"]
        ]

        for col, v in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.border = border
            cell.alignment = Alignment(horizontal="center")

        row += 1

    for col in range(1, 9):
        ws.column_dimensions[chr(64 + col)].width = 20

    wb.save(excel_path)
    print(f"Excel saved: {excel_path}")


def handle_calibration_pair():
    global final_mm_per_px, mode, current_points

    p1, p2 = current_points[-2], current_points[-1]
    px_dist = distance(p1, p2)

    real_mm = simpledialog.askfloat(
        "Calibration",
        "Enter the real distance between these two points in mm:",
        minvalue=0.0001
    )

    if real_mm is None:
        current_points = current_points[:-2]
        return

    mm_per_px = real_mm / px_dist

    calibration_records.append({
        "index": len(calibration_records) + 1,
        "p1": p1,
        "p2": p2,
        "pixel_distance": px_dist,
        "real_mm": real_mm,
        "mm_per_px": mm_per_px
    })

    print(f"\nCalibration {len(calibration_records)} completed")
    print(f"Pixel distance: {px_dist:.3f} px")
    print(f"Real distance: {real_mm:.3f} mm")
    print(f"mm/px: {mm_per_px:.9f}")

    current_points = []

    if len(calibration_records) == 2:
        final_mm_per_px = np.mean([r["mm_per_px"] for r in calibration_records])
        mode = "measurement"

        messagebox.showinfo(
            "Calibration completed",
            f"Two calibrations completed.\n\nMean mm/px = {final_mm_per_px:.9f}\n\nNow enter measurement mode."
        )


def handle_measurement_pair():
    global current_points

    p1, p2 = current_points[-2], current_points[-1]
    px_dist = distance(p1, p2)
    mm_dist = px_dist * final_mm_per_px

    measurement_records.append({
        "index": len(measurement_records) + 1,
        "p1": p1,
        "p2": p2,
        "pixel_distance": px_dist,
        "distance_mm": mm_dist
    })

    print(f"\nMeasurement {len(measurement_records)} completed")
    print(f"Pixel distance: {px_dist:.3f} px")
    print(f"Real distance: {mm_dist:.3f} mm")

    current_points = []
    save_excel()


def mouse_event(event, x, y, flags, param):
    global scale, offset_x, offset_y
    global dragging, last_mouse
    global pending_calibration, pending_measurement

    if event == cv2.EVENT_LBUTTONDOWN:
        ix, iy = screen_to_image(x, y)

        if 0 <= ix < img_w and 0 <= iy < img_h:
            current_points.append((ix, iy))
            print(f"Point selected: ({ix}, {iy})")

            if mode == "calibration" and len(current_points) == 2:
                pending_calibration = True

            elif mode == "measurement" and len(current_points) == 2:
                pending_measurement = True

    elif event == cv2.EVENT_MOUSEWHEEL:
        old_scale = scale

        if flags > 0:
            scale *= 1.15
        else:
            scale /= 1.15

        scale = max(0.1, min(scale, 12.0))

        ix_before = (x - offset_x) / old_scale
        iy_before = (y - offset_y) / old_scale

        offset_x = int(x - ix_before * scale)
        offset_y = int(y - iy_before * scale)

    elif event == cv2.EVENT_RBUTTONDOWN or event == cv2.EVENT_MBUTTONDOWN:
        dragging = True
        last_mouse = (x, y)

    elif event == cv2.EVENT_MOUSEMOVE and dragging:
        dx = x - last_mouse[0]
        dy = y - last_mouse[1]

        offset_x += dx
        offset_y += dy

        last_mouse = (x, y)

    elif event == cv2.EVENT_RBUTTONUP or event == cv2.EVENT_MBUTTONUP:
        dragging = False
        last_mouse = None


def draw():
    canvas = np.ones((canvas_h, canvas_w, 3), dtype=np.uint8) * 235

    resized = cv2.resize(img, None, fx=scale, fy=scale)
    rh, rw = resized.shape[:2]

    x1 = max(offset_x, 0)
    y1 = max(offset_y, 0)
    x2 = min(offset_x + rw, canvas_w)
    y2 = min(offset_y + rh, canvas_h)

    sx1 = max(-offset_x, 0)
    sy1 = max(-offset_y, 0)
    sx2 = sx1 + (x2 - x1)
    sy2 = sy1 + (y2 - y1)

    if x2 > x1 and y2 > y1:
        canvas[y1:y2, x1:x2] = resized[sy1:sy2, sx1:sx2]

    overlay = canvas.copy()

    for r in measurement_records:
        p1 = image_to_screen(r["p1"])
        p2 = image_to_screen(r["p2"])
        cv2.line(overlay, p1, p2, (0, 180, 255), 4)

    canvas = cv2.addWeighted(overlay, 0.45, canvas, 0.55, 0)

    for r in measurement_records:
        p1 = image_to_screen(r["p1"])
        p2 = image_to_screen(r["p2"])
        mx, my = midpoint(p1, p2)

        text = f"{r['distance_mm']:.2f} mm"

        cv2.putText(canvas, text, (mx + 8, my - 8),
                    cv2.FONT_HERSHEY_SIMPLEX, 0.75, (0, 0, 0), 4)
        cv2.putText(canvas, text, (mx + 8, my - 8),
                    cv2.FONT_HERSHEY_SIMPLEX, 0.75, (255, 255, 255), 2)

    for p in current_points:
        sx, sy = image_to_screen(p)

        if mode == "calibration":
            color = (0, 0, 255)
        else:
            color = (0, 255, 0)

        cv2.circle(canvas, (sx, sy), 6, color, -1)

    if len(current_points) == 2:
        p1 = image_to_screen(current_points[0])
        p2 = image_to_screen(current_points[1])
        cv2.line(canvas, p1, p2, (255, 255, 0), 2)

    panel = canvas.copy()
    cv2.rectangle(panel, (15, 15), (535, 205), (255, 255, 255), -1)
    canvas = cv2.addWeighted(panel, 0.75, canvas, 0.25, 0)

    info = [
        f"Mode: {mode.upper()}",
        f"Calibration: {len(calibration_records)} / 2",
        f"Measurements: {len(measurement_records)}",
        "Left click: select point",
        "Mouse wheel: zoom",
        "Right / middle drag: move image",
        "Backspace: delete last point / last measurement",
        "F: fit image to window",
        "ESC: exit"
    ]

    y = 40
    for line in info:
        cv2.putText(canvas, line, (30, y),
                    cv2.FONT_HERSHEY_SIMPLEX, 0.58, (0, 0, 0), 2)
        y += 19

    if final_mm_per_px is not None:
        cv2.putText(canvas, f"Mean mm/px: {final_mm_per_px:.9f}",
                    (30, y + 12),
                    cv2.FONT_HERSHEY_SIMPLEX, 0.58, (0, 120, 0), 2)

    cv2.imshow(WINDOW_NAME, canvas)


cv2.namedWindow(WINDOW_NAME, cv2.WINDOW_NORMAL)
cv2.resizeWindow(WINDOW_NAME, canvas_w, canvas_h)
cv2.setMouseCallback(WINDOW_NAME, mouse_event)

fit_image_to_window()

print("\nProgram started")
print(f"Image: {IMAGE_PATH}")
print(f"Excel output: {excel_path}")

print("""
Workflow:

1. First calibration: click two points on the ruler, then enter the real distance in mm.
2. Second calibration: click another two points on the ruler, then enter the real distance in mm.
3. The program calculates the mean mm/px automatically.
4. Measurement mode starts automatically.
5. Click two points to measure distance.
6. A semi-transparent line and the distance in mm will be displayed.
7. Excel is saved automatically after each measurement.

Shortcuts:
Backspace: delete current point; if no current point, delete the last measurement
F: fit image to window
ESC: exit
""")

while True:
    draw()

    if pending_calibration:
        cv2.waitKey(100)
        pending_calibration = False
        handle_calibration_pair()

    if pending_measurement:
        cv2.waitKey(100)
        pending_measurement = False
        handle_measurement_pair()

    key = cv2.waitKey(20) & 0xFF

    if key == 27:
        break

    elif key == 8:
        if current_points:
            removed = current_points.pop()
            print(f"Deleted current point: {removed}")
        elif measurement_records:
            removed = measurement_records.pop()
            print(f"Deleted last measurement: {removed['index']}")
            save_excel()

    elif key == ord("f"):
        fit_image_to_window()
        print("Image fitted to window.")

cv2.destroyAllWindows()
